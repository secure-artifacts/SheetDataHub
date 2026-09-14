from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from sheet_hub.config_store import ConfigStore
from sheet_hub.database import AggregateDatabase
from sheet_hub.engine import DataEngine, parse_date
from sheet_hub.models import Record, SourceConfig
from sheet_hub.source_reader import SourceReader, canonicalize, choose_sheets, google_retry, parse_schema_lines
from sheet_hub.ui import DEFAULT_QUERY_RESULT_FIELDS, query_result_headers
from sheet_hub.version import APP_VERSION, download_release_installer, fetch_latest_release, is_newer, parse_version


class RuleTests(unittest.TestCase):
    def test_empty_include_runs_all_except_excluded(self):
        selected, missing = choose_sheets(["订单", "统计", "客户"], [], ["统计"])
        self.assertEqual(selected, ["订单", "客户"])
        self.assertEqual(missing, [])

    def test_include_only_and_exclude_wins(self):
        selected, missing = choose_sheets(["订单A", "订单B"], ["订单A", "订单B", "不存在"], ["订单B"])
        self.assertEqual(selected, ["订单A"])
        self.assertEqual(missing, ["不存在"])

    def test_version_compare(self):
        self.assertEqual(parse_version("v1.2.0"), (1, 2, 0))
        self.assertTrue(is_newer("1.2.1", "1.2.0"))
        self.assertFalse(is_newer("1.2.0", "1.2.0"))
        self.assertFalse(is_newer("1.1.9", APP_VERSION))
        self.assertEqual(parse_schema_lines("专页ID\n姓名\n\n号码\n"), ["专页ID", "姓名", "", "号码"])

    def test_space_separated_sheet_exclusions_are_supported(self):
        selected, _ = choose_sheets(["数据", "Index", "清理", "备份链接"], [], ["Index 清理 备份链接"])
        self.assertEqual(selected, ["数据"])

    def test_alias_mapping(self):
        result = canonicalize({"手机号": 13800138000, "渠道": "广告"}, {"号码": ["手机号"], "来源": ["渠道"]})
        self.assertEqual(result["号码"], "13800138000")
        self.assertEqual(result["来源"], "广告")

    def test_date_parsing(self):
        self.assertEqual(parse_date("2026/09/13"), date(2026, 9, 13))
        self.assertEqual(parse_date("2026年9月13日"), date(2026, 9, 13))
        self.assertIsNone(parse_date("not-a-date"))

    def test_update_release_selects_and_downloads_installer(self):
        class FakeJsonResponse:
            status_code = 200

            def raise_for_status(self):
                return None

            def json(self):
                return {
                    "tag_name": "v9.9.9",
                    "html_url": "https://example.test/release",
                    "assets": [{
                        "name": "SheetDataHub-Setup.exe",
                        "browser_download_url": "https://example.test/setup.exe",
                        "size": 1024 * 1024 + 2,
                    }],
                }

        with patch("sheet_hub.version.requests.get", return_value=FakeJsonResponse()):
            info = fetch_latest_release()
        self.assertEqual(info["installer_url"], "https://example.test/setup.exe")

        payload = b"MZ" + (b"x" * (1024 * 1024))

        class FakeDownloadResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def raise_for_status(self):
                return None

            def iter_content(self, chunk_size):
                yield payload

        with tempfile.TemporaryDirectory() as directory, patch(
            "sheet_hub.version.requests.get", return_value=FakeDownloadResponse()
        ):
            path = download_release_installer(info, directory)
            self.assertEqual(path.read_bytes()[:2], b"MZ")

    def test_google_429_is_retried(self):
        class QuotaError(Exception):
            response = SimpleNamespace(status_code=429)

        calls = {"count": 0}

        def operation():
            calls["count"] += 1
            if calls["count"] == 1:
                raise QuotaError()
            return "ok"

        with patch("sheet_hub.source_reader.time.sleep"), patch("sheet_hub.source_reader.random.random", return_value=0):
            self.assertEqual(google_retry(operation), "ok")
        self.assertEqual(calls["count"], 2)

    def test_phone_matching_normalizes_common_formats(self):
        self.assertEqual(DataEngine._match_value("号码", "258-851-758692"), "258851758692")
        self.assertEqual(DataEngine._match_value("号码", "258851758692.0"), "258851758692")
        self.assertEqual(DataEngine._match_value("号码", "2.58851758692E+11"), "258851758692")

    def test_existing_aliases_are_migrated(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            aliases = store.get("field_aliases")
            aliases["号码"] = ["号码", "手机号"]
            store.set("field_aliases", aliases)
            reopened = ConfigStore(directory)
            self.assertIn("手机号码", reopened.get("field_aliases")["号码"])

    def test_query_source_is_remembered(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            self.assertEqual(store.get("query_source"), "extract")
            store.set("query_source", "direct")
            store.set("google_output_url", "https://docs.google.com/spreadsheets/d/abc")
            store.set("google_output_sheet", "提取结果")
            reopened = ConfigStore(directory)
            self.assertEqual(reopened.get("query_source"), "direct")
            self.assertEqual(reopened.get("google_output_sheet"), "提取结果")
            self.assertTrue(reopened.get("query_exact"))
            self.assertFalse(reopened.get("query_fuzzy"))
            self.assertFalse(reopened.get("query_date_enabled"))

    def test_query_result_headers_keep_phone_format(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            source_headers = ["预交表汇总", "见证状态", "交教会日期", "摸底/推广", "线索电话号码"]
            self.assertEqual(
                query_result_headers(store, "extract", "摸底/推广", source_headers),
                DEFAULT_QUERY_RESULT_FIELDS,
            )
            self.assertEqual(
                query_result_headers(store, "direct", "线索电话号码", source_headers),
                DEFAULT_QUERY_RESULT_FIELDS,
            )
            self.assertEqual(
                query_result_headers(store, "direct", "摸底/推广", source_headers),
                source_headers,
            )


class DatabaseTests(unittest.TestCase):
    def test_sharding_and_query(self):
        with tempfile.TemporaryDirectory() as directory:
            db = AggregateDatabase(directory, 1000)
            records = [Record("s", "源", "g", "订单", index, {"号码": str(index)}, str(index)) for index in range(1001)]
            rows, databases = db.replace_all(records)
            self.assertEqual(rows, 1001)
            self.assertEqual(databases, 2)
            self.assertEqual(db.query("号码", "1000")[0].row_number, 1000)

    def test_extract_dedup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            records = [Record("s", "源", "g", "订单", 2, {"日期": "2026-09-13", "号码": "123"}, "hash")]
            engine.database.replace_all(records)
            output1 = root / "one.xlsx"
            output2 = root / "two.xlsx"
            first = engine.extract("日期", date(2026, 9, 1), date(2026, 9, 30), output1, ["号码", "日期"])
            second = engine.extract("日期", date(2026, 9, 1), date(2026, 9, 30), output2, ["号码", "日期"])
            self.assertEqual(first["written"], 1)
            self.assertEqual(second["written"], 0)
            self.assertTrue(output2.exists())

    def test_batch_query_keeps_missing_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1008-李薇", 2, {"号码": "123", "名字": "李薇", "日期": "2026-09-13"}, "h1")
            ])
            results = engine.query_many("号码", ["123", "999"])
            self.assertEqual(results[0][1].sheet_name, "1008-李薇")
            self.assertIsNone(results[1][1])

    def test_fuzzy_source_query_matches_original_and_corrected_sheet_name(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1233-赵刚", 2, {"号码": "123"}, "h1")
            ])
            for value in ("赵刚", "1233", "1233-赵刚", "赵刚-1233-专页后台"):
                results = engine.query_many("来源", [value], exact=False)
                self.assertEqual(results[0][1].sheet_name, "1233-赵刚")

    def test_exact_source_query_matches_corrected_sheet_name(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1233-赵刚", 2, {"号码": "123"}, "h1")
            ])
            results = engine.query_many("来源", ["赵刚-1233-专页后台"], exact=True)
            self.assertEqual(results[0][1].sheet_name, "1233-赵刚")

    def test_query_can_limit_results_by_date_range(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "A", 2, {"名字": "刘海", "日期": "2026-09-01"}, "h1"),
                Record("s", "源", "g", "B", 3, {"名字": "刘海", "日期": "2026-09-13"}, "h2"),
                Record("s", "源", "g", "C", 4, {"名字": "刘海", "日期": "无效日期"}, "h3"),
            ])
            results = engine.query_many(
                "名字", ["刘海"], exact=True, date_field="日期",
                start_date=date(2026, 9, 10), end_date=date(2026, 9, 20),
            )
            self.assertEqual(len(results), 1)
            self.assertEqual(results[0][1].sheet_name, "B")

    def test_custom_headers_are_resolved_for_query_date_and_phone(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "A", 2,
                    {
                        "交教会日期": "2026年9月13日",
                        "线索电话号码": "258-851-758692",
                        "摸底/推广": "简\u200b 单",
                    },
                    "h1",
                ),
            ])
            results = engine.query_many(
                "摸底/推广", ["简单"], exact=False, date_field="日期",
                start_date=date(2026, 9, 1), end_date=date(2026, 9, 30),
            )
            self.assertIsNotNone(results[0][1])
            self.assertEqual(engine.query("号码", "258851758692")[0].sheet_name, "A")
            self.assertEqual(
                engine.suggest_field(["交教会日期", "摸底/推广"], "日期"),
                "交教会日期",
            )

    def test_extract_resolves_custom_date_and_dedup_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "A", 2,
                    {"交教会日期": "2026-09-13", "线索电话号码": "123"},
                    "h1",
                ),
            ])
            first = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                root / "one.xlsx", ["号码", "日期"],
            )
            second = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                root / "two.xlsx", ["号码", "日期"],
            )
            self.assertEqual(first["written"], 1)
            self.assertEqual(second["duplicates"], 1)

    def test_extract_output_schema_supports_column_mapping_objects(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("extract_column_schema_enabled", True)
            store.set("extract_column_schema", [
                {"name": "线索电话号码", "column": "S", "enabled": True},
                {"name": "交教会日期", "column": "F", "enabled": True},
            ])
            engine = DataEngine(store)
            headers = engine._preferred_export_headers([])
            self.assertEqual(headers, ["来源", "线索电话号码", "交教会日期"])
            rows = engine._rows_for_headers(
                [Record("s", "源", "g", "A", 2, {"号码": "123", "日期": "2026-09-13"}, "h")],
                headers,
                store.get("field_aliases"),
            )
            self.assertEqual(rows[0], ["A", "123", "2026-09-13"])

    def test_query_extract_table_is_faster_subset(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            extract_path = Path(directory) / "extract.xlsx"
            DataEngine._write_xlsx(
                extract_path,
                [
                    Record("s", "源", "g", "1008-李薇", 2, {"号码": "123", "名字": "李薇", "日期": "2026-09-13"}, "h1"),
                    Record("s", "源", "g", "2002-王强", 3, {"号码": "456", "名字": "王强", "日期": "2026-09-12"}, "h2"),
                ],
                "提取结果",
            )
            other = openpyxl.load_workbook(extract_path)
            extra = other.create_sheet("其他页")
            extra.append(["来源", "号码"])
            extra.append(["不该读", "999"])
            other.save(extract_path)
            other.close()
            results = engine.query_many(
                "号码",
                ["123", "999"],
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
            )
            self.assertEqual(results[0][1].sheet_name, "1008-李薇")
            self.assertEqual(results[0][1].values["号码"], "123")
            self.assertIsNone(results[1][1])
            found = engine.query(
                "号码",
                "456",
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
            )
            self.assertEqual(found[0].values["名字"], "王强")
            headers = engine.list_query_fields(
                "extract", str(extract_path), "提取结果",
            )
            self.assertIn("来源", headers)
            self.assertTrue(any(name in {"号码", "手机号码"} for name in headers))
            by_header = engine.query(
                "名字",
                "李薇",
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
            )
            self.assertEqual(by_header[0].values["号码"], "123")

    def test_extract_uses_configured_sheet_and_source_first(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            records = [Record("s", "源", "g", "渠道Sheet", 2, {"号码": "123", "日期": "2026-09-13"}, "h")]
            DataEngine._write_xlsx(path, records, "指定工作表")
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(book.sheetnames, ["指定工作表"])
            rows = list(book["指定工作表"].iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "来源")
            self.assertEqual(rows[1][0], "渠道Sheet")
            book.close()

    def test_export_header_aliases_and_order_are_compatible(self):
        aliases = {
            "名字": ["姓名"],
            "号码": ["手机号码"],
            "日期": ["日期"],
        }
        existing = ["来源", "专页ID", "姓名", "手机号码", "日期"]
        required = ["来源", "日期", "名字", "号码", "专页ID"]
        self.assertTrue(DataEngine._headers_compatible(existing, required, aliases))
        records = [Record(
            "s", "源", "g", "1008-李薇", 2,
            {"专页ID": "42", "名字": "李薇", "号码": "258851758692", "日期": "2026-09-13"},
            "h",
        )]
        self.assertEqual(
            DataEngine._rows_for_headers(records, existing, aliases)[0],
            ["1008-李薇", "42", "李薇", "258851758692", "2026-09-13"],
        )


class WorkbookTests(unittest.TestCase):
    def test_local_workbook_source(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "订单"
            sheet.append(["日期", "手机号", "渠道"])
            sheet.append(["2026-09-13", "10086", "搜索"])
            stats = book.create_sheet("统计")
            stats.append(["日期", "手机号"])
            stats.append(["2026-09-13", "should_skip"])
            book.save(path)
            reader = SourceReader({"号码": ["手机号"], "来源": ["渠道"], "日期": ["日期"]}, ["统计"])
            source = SourceConfig("id", "测试", str(path))
            records = reader.read(source)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].values["号码"], "10086")

    def test_sources_can_use_different_column_counts(self):
        with tempfile.TemporaryDirectory() as directory:
            short_path = Path(directory) / "short.xlsx"
            long_path = Path(directory) / "long.xlsx"
            short = openpyxl.Workbook()
            short.active.title = "短表"
            short.active.append(["a", "b", "ignore"])
            short.active.append(["page-1", "13800138000", "skip"])
            short.save(short_path)
            long = openpyxl.Workbook()
            long.active.title = "长表"
            long.active.append(["a", "b", "c", "d"])
            long.active.append(["page-2", "李薇", "https://example.test", "258851758692"])
            long.save(long_path)
            store = ConfigStore(Path(directory) / "config")
            store.save_source(SourceConfig(
                "short", "短表", str(short_path),
                column_schema_enabled=True,
                column_schema=["专页ID", "手机号码"],
            ))
            store.save_source(SourceConfig(
                "long", "长表", str(long_path),
                column_schema_enabled=True,
                column_schema=["专页ID", "姓名", "评论贴文", "手机号码"],
            ))
            engine = DataEngine(store)
            records = engine.read_sources()
            by_source = {record.source_id: record for record in records}
            self.assertEqual(by_source["short"].values["专页ID"], "page-1")
            self.assertEqual(by_source["short"].values["号码"], "13800138000")
            self.assertNotIn("ignore", by_source["short"].values)
            self.assertEqual(by_source["long"].values["名字"], "李薇")
            self.assertEqual(by_source["long"].values["号码"], "258851758692")
            self.assertEqual(by_source["long"].values["评论贴文"], "https://example.test")
            short_fields = engine.list_query_fields("direct", source_id="short")
            long_fields = engine.list_query_fields("direct", source_id="long")
            self.assertIn("手机号码", short_fields)
            self.assertNotIn("评论贴文", short_fields)
            self.assertIn("评论贴文", long_fields)
            self.assertEqual(
                engine.query("手机号码", "13800138000", source="direct", source_id="short")[0].source_id,
                "short",
            )

    def test_manual_column_count_and_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "数据"
            sheet.append(["原表头A", "原表头B", "不读取的列"])
            sheet.append(["page-1", "13800138000", "ignore-me"])
            book.save(path)
            reader = SourceReader(
                {"号码": ["手机号码"]},
                [],
                column_schema=["专页ID", "手机号码"],
            )
            records = reader.read(SourceConfig("id", "手动列", str(path)))
            self.assertEqual(records[0].values["专页ID"], "page-1")
            self.assertEqual(records[0].values["号码"], "13800138000")
            self.assertNotIn("不读取的列", records[0].values)

    def test_column_letter_mapping_skips_leading_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "letters.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "数据"
            sheet.append(["A列", "B列", "C列", "贴文ID", "手机号码", "多余"])
            sheet.append(["x", "y", "z", "page-9", "13800138000", "no"])
            book.save(path)
            reader = SourceReader(
                {"号码": ["手机号码"]},
                [],
                column_schema=[
                    {"name": "贴文ID", "column": "D", "enabled": True},
                    {"name": "手机号码", "column": "E", "enabled": True},
                    {"name": "忽略", "column": "F", "enabled": False},
                ],
            )
            records = reader.read(SourceConfig("id", "列映射", str(path)))
            self.assertEqual(records[0].values["贴文ID"], "page-9")
            self.assertEqual(records[0].values["号码"], "13800138000")
            self.assertNotIn("忽略", records[0].values)
            self.assertNotIn("A列", records[0].values)

    def test_private_source_uses_one_batch_request(self):
        class FakeSpreadsheet:
            def __init__(self):
                self.batch_calls = 0

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [
                    {"properties": {"title": "A"}},
                    {"properties": {"title": "B"}},
                ]}

            def values_batch_get(self, key, ranges, params=None):
                self.batch_calls += 1
                return {"valueRanges": [
                    {"values": [["手机号"], ["111"]]},
                    {"values": [["手机号"], ["222"]]},
                ]}

        spreadsheet = FakeSpreadsheet()
        client = SimpleNamespace(http_client=spreadsheet)
        source = SourceConfig("id", "私有表", "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz", credential_path="fake.json")
        reader = SourceReader({"号码": ["手机号"]}, [])
        with patch.object(SourceReader, "_gspread_client", return_value=client):
            records = reader.read(source)
        self.assertEqual(spreadsheet.batch_calls, 1)
        self.assertEqual([record.values["号码"] for record in records], ["111", "222"])

    def test_google_output_appends_header_and_rows(self):
        class FakeHttp:
            def __init__(self):
                self.appended = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "目标页"}}]}

            def values_get(self, key, range_name, params=None):
                return {"values": []}

            def values_append(self, key, range_name, params=None, body=None):
                self.appended.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            engine = DataEngine(store)
            records = [Record("s", "源", "g", "渠道A", 2, {"号码": "123"}, "h")]
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                engine._write_google_sheet(
                    "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    "目标页",
                    records,
                    "测试",
                )
        self.assertEqual(http.appended[0], ["来源", "号码"])
        self.assertEqual(http.appended[1], ["渠道A", "123"])

    def test_google_output_uses_existing_alias_header_order(self):
        existing = ["来源", "专页ID", "姓名", "标签", "订阅时间", "性别", "评论贴文", "手机号码", "日期"]

        class FakeHttp:
            def __init__(self):
                self.appended = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "测试"}}]}

            def values_get(self, key, range_name, params=None):
                return {"values": [existing]}

            def values_append(self, key, range_name, params=None, body=None):
                self.appended.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            store.set("column_schema_enabled", True)
            store.set("column_schema", existing[1:])
            engine = DataEngine(store)
            records = [Record(
                "s", "源", "g", "1008-李薇", 2,
                {
                    "专页ID": "42", "名字": "李薇", "标签": "A", "订阅时间": "12:00",
                    "性别": "女", "评论贴文": "https://example.test", "号码": "258851758692",
                    "日期": "2026-09-13",
                },
                "h",
            )]
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                engine._write_google_sheet(
                    "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    "测试",
                    records,
                    "测试",
                )
        self.assertEqual(http.appended[0], [
            "1008-李薇", "42", "李薇", "A", "12:00", "女", "https://example.test",
            "258851758692", "2026-09-13",
        ])


if __name__ == "__main__":
    unittest.main()
