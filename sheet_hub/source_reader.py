from __future__ import annotations

import hashlib
import io
import json
import random
import re
import time
from pathlib import Path
from typing import Callable, Iterable

import openpyxl
import requests

from .models import Record, SourceConfig


LogFn = Callable[[str, str], None]


def google_retry(call, logger: LogFn | None = None, attempts: int = 7):
    """Retry Google API quota and transient server errors with bounded backoff."""
    log = logger or (lambda level, message: None)
    for attempt in range(attempts):
        try:
            return call()
        except Exception as exc:
            response = getattr(exc, "response", None)
            status = getattr(response, "status_code", None)
            if status not in {429, 500, 502, 503, 504} or attempt == attempts - 1:
                raise
            wait_seconds = min(45.0, (2 ** attempt) + random.random())
            log("WARNING", f"Google API 暂时限流（{status}），{wait_seconds:.1f} 秒后自动重试 {attempt + 2}/{attempts}")
            time.sleep(wait_seconds)


def spreadsheet_id(url: str) -> str:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url)
    if match:
        return match.group(1)
    if re.fullmatch(r"[a-zA-Z0-9_-]{20,}", url.strip()):
        return url.strip()
    return ""


def split_names(text: str | Iterable[str]) -> list[str]:
    if not isinstance(text, str):
        return [str(item).strip() for item in text if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，;；\n]+", text) if item.strip()]


def choose_sheets(all_names: list[str], include: list[str], exclude: list[str]) -> tuple[list[str], list[str]]:
    available_names = {name.casefold() for name in all_names}

    def expand(configured: list[str]) -> list[str]:
        expanded: list[str] = []
        for name in configured:
            clean = name.strip()
            if clean.casefold() in available_names:
                expanded.append(clean)
            else:
                expanded.extend(part for part in re.split(r"\s+", clean) if part)
        return expanded

    include = expand(include)
    exclude = expand(exclude)
    excluded = {name.casefold() for name in exclude}
    available = {name.casefold(): name for name in all_names}
    requested = include if include else all_names
    selected: list[str] = []
    missing: list[str] = []
    for name in requested:
        actual = available.get(name.casefold())
        if actual is None:
            missing.append(name)
        elif actual.casefold() not in excluded and actual not in selected:
            selected.append(actual)
    return selected, missing


def normalize_headers(values: list[object]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"未命名列{index}"
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def canonicalize(row: dict[str, object], aliases: dict[str, list[str]]) -> dict[str, str]:
    lookup = {str(key).strip().casefold(): value for key, value in row.items()}
    result: dict[str, str] = {}
    used: set[str] = set()
    for canonical, candidates in aliases.items():
        for candidate in [canonical, *candidates]:
            key = str(candidate).strip().casefold()
            if key in lookup:
                value = lookup[key]
                result[canonical] = "" if value is None else str(value).strip()
                used.add(key)
                break
    for header, value in row.items():
        if str(header).strip().casefold() not in used:
            result[str(header)] = "" if value is None else str(value).strip()
    return result


class SourceReader:
    def __init__(
        self,
        aliases: dict[str, list[str]],
        global_excludes: list[str],
        logger: LogFn | None = None,
        column_schema: list[str] | None = None,
    ):
        self.aliases = aliases
        self.global_excludes = global_excludes
        self.logger = logger or (lambda level, message: None)
        self.column_schema = [str(name).strip() for name in (column_schema or [])]

    def _headers(self, actual_values: list[object]) -> tuple[list[str], int | None]:
        if not self.column_schema:
            return normalize_headers(actual_values), None
        count = len(self.column_schema)
        padded = list(actual_values[:count]) + [""] * max(0, count - len(actual_values))
        configured = [self.column_schema[index] or padded[index] for index in range(count)]
        return normalize_headers(configured), count

    def read(self, source: SourceConfig) -> list[Record]:
        sid = spreadsheet_id(source.url)
        if not sid:
            path = Path(source.url)
            if path.exists() and path.suffix.lower() in {".xlsx", ".xlsm"}:
                return self._read_workbook(source, path.read_bytes(), path.stem)
            raise ValueError("无法识别 Google 表格链接或本地 Excel 文件")
        credential = source.credential_path.strip()
        if credential:
            return self._read_private(source, sid, credential)
        return self._read_public(source, sid)

    def list_sheets(self, source: SourceConfig) -> list[str]:
        sid = spreadsheet_id(source.url)
        if source.credential_path.strip():
            client = self._gspread_client(source.credential_path)
            metadata = google_retry(
                lambda: client.http_client.fetch_sheet_metadata(
                    sid, params={"fields": "sheets.properties.title"}
                ),
                self.logger,
            )
            return [item["properties"]["title"] for item in metadata.get("sheets", [])]
        response = requests.get(
            f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx",
            timeout=60,
        )
        response.raise_for_status()
        workbook = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        return names

    def _read_public(self, source: SourceConfig, sid: str) -> list[Record]:
        self.logger("INFO", f"下载公开表格：{source.name}")
        response = requests.get(
            f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx",
            timeout=120,
        )
        if response.status_code in {401, 403}:
            raise PermissionError("表格不可公开读取，请共享为可查看或配置服务账号 JSON")
        response.raise_for_status()
        return self._read_workbook(source, response.content, sid)

    def _read_workbook(self, source: SourceConfig, content: bytes, sid: str) -> list[Record]:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        excludes = [*self.global_excludes, *source.exclude_sheets]
        selected, missing = choose_sheets(workbook.sheetnames, source.include_sheets, excludes)
        for name in missing:
            self.logger("WARNING", f"{source.name}：指定的 Sheet 不存在：{name}")
        records: list[Record] = []
        for sheet_name in selected:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            for _ in range(max(0, source.header_row - 1)):
                next(rows, None)
            header_values = next(rows, None)
            if not header_values:
                self.logger("WARNING", f"{source.name}/{sheet_name}：空 Sheet，已跳过")
                continue
            headers, column_count = self._headers(list(header_values))
            count = 0
            for row_number, values in enumerate(rows, start=source.header_row + 1):
                if not values or not any(value not in (None, "") for value in values):
                    continue
                row_values = values[:column_count] if column_count else values
                raw = dict(zip(headers, row_values))
                mapped = canonicalize(raw, self.aliases)
                digest = self._hash(sid, sheet_name, row_number, mapped)
                records.append(Record(source.id, source.name, sid, sheet_name, row_number, mapped, digest))
                count += 1
            self.logger("INFO", f"{source.name}/{sheet_name}：读取 {count} 行")
        workbook.close()
        return records

    def _read_private(self, source: SourceConfig, sid: str, credential: str) -> list[Record]:
        client = self._gspread_client(credential)
        metadata = google_retry(
            lambda: client.http_client.fetch_sheet_metadata(
                sid, params={"fields": "sheets.properties(title,hidden)"}
            ),
            self.logger,
        )
        all_names = [item["properties"]["title"] for item in metadata.get("sheets", [])]
        excludes = [*self.global_excludes, *source.exclude_sheets]
        selected, missing = choose_sheets(all_names, source.include_sheets, excludes)
        for name in missing:
            self.logger("WARNING", f"{source.name}：指定的 Sheet 不存在：{name}")
        records: list[Record] = []
        values_by_sheet: dict[str, list[list[str]]] = {}
        for start in range(0, len(selected), 100):
            names = selected[start:start + 100]
            ranges = [f"'{name.replace(chr(39), chr(39) * 2)}'!A:ZZZ" for name in names]
            response = google_retry(
                lambda ranges=ranges: client.http_client.values_batch_get(
                    sid,
                    ranges,
                    params={"majorDimension": "ROWS", "valueRenderOption": "FORMATTED_VALUE"},
                ),
                self.logger,
            )
            returned = response.get("valueRanges", [])
            for index, sheet_name in enumerate(names):
                values_by_sheet[sheet_name] = returned[index].get("values", []) if index < len(returned) else []
            self.logger("INFO", f"{source.name}：已批量读取 {len(names)} 个子 Sheet（1 次 API 请求）")
        for sheet_name in selected:
            values = values_by_sheet.get(sheet_name, [])
            if len(values) < source.header_row:
                self.logger("WARNING", f"{source.name}/{sheet_name}：空 Sheet，已跳过")
                continue
            headers, column_count = self._headers(values[source.header_row - 1])
            count = 0
            for row_number, row_values in enumerate(values[source.header_row:], start=source.header_row + 1):
                if not any(str(value).strip() for value in row_values):
                    continue
                limited_values = row_values[:column_count] if column_count else row_values
                raw = dict(zip(headers, limited_values))
                mapped = canonicalize(raw, self.aliases)
                digest = self._hash(sid, sheet_name, row_number, mapped)
                records.append(Record(source.id, source.name, sid, sheet_name, row_number, mapped, digest))
                count += 1
            self.logger("INFO", f"{source.name}/{sheet_name}：读取 {count} 行")
        return records

    @staticmethod
    def _gspread_client(credential: str):
        import gspread

        path = Path(credential)
        if not path.exists():
            raise FileNotFoundError(f"服务账号文件不存在：{credential}")
        return gspread.service_account(filename=str(path))

    @staticmethod
    def _hash(sid: str, sheet: str, row_number: int, values: dict[str, str]) -> str:
        raw = json.dumps([sid, sheet, row_number, values], ensure_ascii=False, sort_keys=True)
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()
