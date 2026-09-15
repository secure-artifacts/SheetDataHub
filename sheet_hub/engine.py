from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter
from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import openpyxl

from .config_store import ConfigStore
from .database import AggregateDatabase, SourceCache
from .models import Record, SourceConfig
from .source_reader import SourceReader, column_index, google_retry, schema_field_names, spreadsheet_id


ProgressFn = Callable[[str], None]


def parse_date(value: str) -> date | None:
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"\s*(年|月)\s*", lambda match: "-" if match.group(1) in {"年", "月"} else match.group(0), text)
    text = re.sub(r"\s*日\s*$", "", text)
    if "T" in text:
        text = text.split("T", 1)[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


class DataEngine:
    SEMANTIC_FIELD_HINTS = {
        "日期": ("日期", "date", "datetime"),
        "号码": ("手机", "电话", "号码", "phone", "number", "tel"),
        "名字": ("姓名", "名字", "名称", "name"),
        "链接": ("链接", "网址", "link", "url"),
        "来源": ("来源", "渠道", "平台", "source"),
    }

    def __init__(self, store: ConfigStore, progress: ProgressFn | None = None):
        self.store = store
        self.progress = progress or (lambda message: None)

    @property
    def database(self) -> AggregateDatabase:
        return AggregateDatabase(
            self.store.data_dir / "databases",
            int(self.store.get("max_rows_per_db", 500000)),
        )

    @property
    def cache(self) -> SourceCache:
        return SourceCache(self.store.data_dir / "source_cache")

    def _log(self, operation: str, level: str, message: str, detail: str = "") -> None:
        self.store.log(level, operation, message, detail)
        self.progress(message)

    def _schema_for_source(self, source: SourceConfig | None, use_schema: bool) -> list[str]:
        if not use_schema:
            return []
        if source is not None and source.column_schema_enabled and source.column_schema:
            return list(source.column_schema)
        if self.store.get("column_schema_enabled", False):
            return list(self.store.get("column_schema", []) or [])
        return []

    def _reader(
        self,
        operation: str,
        *,
        source: SourceConfig | None = None,
        use_schema: bool = True,
        schema: list[str] | None = None,
        global_excludes: list[str] | None = None,
    ) -> SourceReader:
        resolved = schema if schema is not None else self._schema_for_source(source, use_schema)
        excludes = self.store.get("global_excludes", []) if global_excludes is None else global_excludes
        return SourceReader(
            self.store.get("field_aliases", {}),
            excludes,
            lambda level, message: self._log(operation, level, message),
            resolved,
        )

    def read_sources(self, operation: str = "读取数据源", source_id: str = "") -> list[Record]:
        sources = [source for source in self.store.load_sources() if source.enabled]
        if source_id.strip():
            sources = [source for source in sources if source.id == source_id.strip()]
        if not sources:
            raise ValueError("没有启用的数据源" if not source_id.strip() else "没有找到所选数据源")
        records: list[Record] = []
        failures = 0
        for source in sources:
            try:
                if not source.credential_path:
                    source.credential_path = str(self.store.get("credential_path", "")).strip()
                self._log(operation, "INFO", f"开始读取：{source.name}")
                loaded = self._reader(operation, source=source).read(source)
                records.extend(loaded)
                self._log(operation, "INFO", f"完成读取：{source.name}，共 {len(loaded)} 行")
            except Exception as exc:
                failures += 1
                self._log(operation, "ERROR", f"读取失败：{source.name}", repr(exc))
        if failures == len(sources):
            raise RuntimeError("所有数据源均读取失败，请查看日志")
        return records

    def _headers_for_source(self, source: SourceConfig, records: list[Record]) -> list[str]:
        if source.column_schema_enabled:
            names = schema_field_names(source.column_schema)
            if names:
                return names
        seen: list[str] = []
        for record in records[:50]:
            for key in record.values:
                if key not in seen:
                    seen.append(key)
        return seen

    def refresh_cache(self, source_ids: list[str] | None = None, include_extract: bool = False) -> dict[str, int]:
        operation = "刷新缓存"
        selected = [source for source in self.store.load_sources() if source.enabled]
        if source_ids:
            wanted = {item.strip() for item in source_ids if item.strip()}
            selected = [source for source in selected if source.id in wanted]
        if not selected and not include_extract:
            raise ValueError("没有可刷新的数据源")
        total = 0
        for source in selected:
            self._log(operation, "INFO", f"正在缓存：{source.name}")
            records = self.read_sources(operation, source.id)
            total += self.cache.replace(source.id, records, self._headers_for_source(source, records), source.url)
        if include_extract:
            dest_type = str(self.store.get("extract_destination_type", "local") or "local")
            target = str(self.store.get("google_output_url") if dest_type == "google" else self.store.get("extract_output_path") or "")
            sheet = str(self.store.get("google_output_sheet") or "提取结果")
            if target.strip():
                records = self.read_extract_table(target, sheet)
                total += self.cache.replace(
                    "extract-table",
                    records,
                    schema_field_names(self.store.get("extract_column_schema", [])) or None,
                    f"{target}#{sheet}",
                )
        self._log(operation, "INFO", f"缓存更新完成：{total} 行")
        return {"rows": total, "sources": len(selected)}

    def cached_records(self, source_ids: list[str], refresh: bool = False) -> list[Record]:
        missing = [source_id for source_id in source_ids if refresh or not self.cache.has(source_id)]
        if missing:
            self.refresh_cache(missing)
        records: list[Record] = []
        for source_id in source_ids:
            records.extend(self.cache.load(source_id))
        return records

    def sync(
        self,
        source_ids: list[str] | None = None,
        write_local_db: bool | None = None,
        google_output_url: str = "",
        google_output_sheet: str = "汇总结果",
        local_xlsx: str = "",
    ) -> dict[str, int]:
        operation = "汇总同步"
        self._log(operation, "INFO", "汇总任务开始")
        selected = [source for source in self.store.load_sources() if source.enabled]
        if source_ids:
            wanted = {item.strip() for item in source_ids if item.strip()}
            selected = [source for source in selected if source.id in wanted]
        if not selected:
            raise ValueError("请选择至少一个数据源")
        records: list[Record] = []
        for source in selected:
            loaded = self.read_sources(operation, source.id)
            self.cache.replace(source.id, loaded, self._headers_for_source(source, loaded), source.url)
            records.extend(loaded)
            self._log(operation, "INFO", f"已缓存 {source.name}：{len(loaded)} 行")
        unique = list({record.row_hash: record for record in records}.values())
        write_db = self.store.get("write_aggregate", True) if write_local_db is None else write_local_db
        db_count = 0
        if write_db:
            row_count, db_count = self.database.replace_all(unique)
        else:
            row_count = len(unique)
        if google_output_url.strip():
            self._write_google_sheet(google_output_url, google_output_sheet or "汇总结果", unique, operation)
            self._log(operation, "INFO", f"已写入 Google 表格工作表“{google_output_sheet or '汇总结果'}”")
        if local_xlsx.strip():
            path = Path(local_xlsx)
            path.parent.mkdir(parents=True, exist_ok=True)
            headers = self._preferred_export_headers(unique)
            self._write_xlsx(path, unique, google_output_sheet or "汇总结果", headers, self.store.get("field_aliases", {}))
            self._log(operation, "INFO", f"已写入本地文件：{path}")
        self._log(operation, "INFO", f"汇总完成：{row_count} 行，缓存 {len(selected)} 个数据源")
        return {"rows": row_count, "databases": db_count, "sources": len(selected)}

    def read_extract_table(self, target: str, sheet_name: str = "") -> list[Record]:
        path_or_url = str(target or "").strip()
        if not path_or_url:
            raise ValueError("请先在「时间提取」页填写提取表的目标表格链接或输出文件")
        sheet = str(sheet_name or "").strip() or "提取结果"
        operation = "查询提取表"
        self._log(operation, "INFO", f"读取提取表：{path_or_url} / {sheet}")
        source = SourceConfig(
            id="extract-table",
            name="提取表",
            url=path_or_url,
            include_sheets=[sheet],
            credential_path=str(self.store.get("credential_path", "")).strip(),
        )
        extract_schema = list(self.store.get("extract_column_schema", []) or [])
        use_extract_schema = bool(self.store.get("extract_column_schema_enabled", False) and schema_field_names(extract_schema))
        records = self._reader(
            operation,
            use_schema=use_extract_schema,
            schema=extract_schema if use_extract_schema else None,
            global_excludes=[],
        ).read(source)
        for record in records:
            origin = str(record.values.get("来源", "")).strip()
            if origin:
                record.sheet_name = origin
        self._log(operation, "INFO", f"提取表读取完成：{len(records)} 行")
        return records

    @staticmethod
    def query_source_mode(source: str | None, direct: bool) -> str:
        mode = str(source or "").strip()
        if mode in {"aggregate", "extract", "direct"}:
            return mode
        return "direct" if direct else "aggregate"

    @staticmethod
    def _unique_headers(items: list[str]) -> list[str]:
        seen: set[str] = set()
        result: list[str] = []
        for item in items:
            name = str(item).strip()
            key = name.casefold()
            if name and key not in seen:
                seen.add(key)
                result.append(name)
        return result

    def _canonical_field(self, field: str) -> str:
        folded = str(field).strip().casefold()
        if not folded:
            return field
        for canonical, names in self.store.get("field_aliases", {}).items():
            options = [canonical, *names]
            if folded in {str(name).strip().casefold() for name in options}:
                return str(canonical)
        return str(field).strip()

    def _field_kind(self, field: str) -> str:
        canonical = self._canonical_field(field)
        if canonical in self.SEMANTIC_FIELD_HINTS:
            return canonical
        folded = str(field).strip().casefold()
        matches = [
            standard
            for standard, hints in self.SEMANTIC_FIELD_HINTS.items()
            if any(hint in folded for hint in hints)
        ]
        return matches[0] if len(matches) == 1 else canonical

    def _field_value(self, record: Record, field: str) -> str:
        folded = str(field).strip().casefold()
        aliases = self.store.get("field_aliases", {})
        candidates = [str(field).strip()]
        for canonical, names in aliases.items():
            options = [str(canonical), *[str(name) for name in names]]
            if folded in {name.strip().casefold() for name in options if name.strip()}:
                candidates.extend(options)
                break
        seen: set[str] = set()
        for candidate in candidates:
            key = candidate.strip().casefold()
            if not key or key in seen:
                continue
            seen.add(key)
            direct = record.values.get(candidate, "")
            if direct:
                return str(direct)
            for stored, value in record.values.items():
                if str(stored).strip().casefold() == key and str(value).strip():
                    return str(value)
        canonical = self._field_kind(field)
        hints = self.SEMANTIC_FIELD_HINTS.get(canonical, ())
        semantic_matches = [
            value
            for stored, value in record.values.items()
            if str(value).strip() and any(hint in str(stored).strip().casefold() for hint in hints)
        ]
        # Only infer a custom field when it is unambiguous. This supports names
        # such as “交教会日期” and “线索电话号码” without guessing between two dates.
        if len(semantic_matches) == 1:
            return str(semantic_matches[0])
        return ""

    def suggest_field(self, fields: list[str], preferred: str) -> str:
        """Find the best visible field for a generic/saved field name."""
        clean_fields = self._unique_headers(fields)
        folded = str(preferred).strip().casefold()
        for field in clean_fields:
            if field.casefold() == folded:
                return field
        canonical = self._field_kind(preferred)
        alias_matches = [field for field in clean_fields if self._field_kind(field) == canonical]
        if len(alias_matches) == 1:
            return alias_matches[0]
        hints = self.SEMANTIC_FIELD_HINTS.get(canonical, ())
        semantic = [field for field in clean_fields if any(hint in field.casefold() for hint in hints)]
        return semantic[0] if semantic else ""

    @staticmethod
    def corrected_source(source: str) -> str:
        if "-" not in source:
            return source
        prefix, remainder = source.split("-", 1)
        return f"{remainder}-{prefix}-专页后台"

    def _query_values(self, record: Record, field: str) -> list[str]:
        values = [self._field_value(record, field)]
        if self._field_kind(field).casefold() == "来源":
            values.extend((record.sheet_name, self.corrected_source(record.sheet_name)))
        return self._unique_headers(values)

    def query_display_value(self, record: Record, field: str) -> str:
        folded = str(field).strip().casefold()
        if folded == "修正格式":
            return self.corrected_source(record.sheet_name)
        if self._field_kind(field).casefold() == "来源":
            return record.sheet_name
        return self._field_value(record, field)

    def _peek_headers(self, target: str, sheet_name: str = "") -> list[str]:
        path = Path(str(target or "").strip())
        if not path.exists() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            return []
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            names = workbook.sheetnames
            if not names:
                return []
            title = sheet_name.strip() if sheet_name.strip() in names else names[0]
            row = next(workbook[title].iter_rows(values_only=True), None)
            return [str(cell).strip() for cell in (row or []) if str(cell or "").strip()]
        finally:
            workbook.close()

    def list_query_fields(
        self,
        mode: str = "extract",
        extract_target: str = "",
        extract_sheet: str = "",
        source_id: str = "",
    ) -> list[str]:
        mode = self.query_source_mode(mode, False)
        if mode == "extract":
            schema_names = schema_field_names(self.store.get("extract_column_schema", []))
            if self.store.get("extract_column_schema_enabled", False) and schema_names:
                return self._unique_headers(schema_names)
            cached = self.cache.headers("extract-table")
            if cached:
                return self._unique_headers(cached)
            peeked = self._peek_headers(extract_target, extract_sheet)
            if peeked:
                return self._unique_headers(peeked)
            return []
        if mode == "direct":
            sources = [source for source in self.store.load_sources() if source.enabled]
            if source_id.strip():
                sources = [source for source in sources if source.id == source_id.strip()]
            headers: list[str] = []
            for source in sources:
                if source.column_schema_enabled and source.column_schema:
                    headers.extend(schema_field_names(source.column_schema))
                    continue
                cached = self.cache.headers(source.id)
                if cached:
                    headers.extend(cached)
                    continue
                peeked = self._peek_headers(source.url, (source.include_sheets or [""])[0])
                if peeked:
                    headers.extend(peeked)
            return self._unique_headers(headers)
        headers = []
        for record in self.database.sample_records(20):
            headers.extend(record.values.keys())
        return self._unique_headers(headers)

    def query(
        self,
        field: str,
        value: str,
        direct: bool = False,
        exact: bool = True,
        source: str | None = None,
        extract_target: str = "",
        extract_sheet: str = "",
        source_id: str = "",
        date_field: str = "",
        start_date: date | None = None,
        end_date: date | None = None,
        refresh_cache: bool = False,
    ) -> list[Record]:
        results = self.query_many(
            field, [value], direct, exact, source, extract_target, extract_sheet, source_id,
            date_field, start_date, end_date, refresh_cache,
        )
        return [record for _, record in results if record is not None]

    def query_many(
        self,
        field: str,
        values: list[str],
        direct: bool = False,
        exact: bool = True,
        source: str | None = None,
        extract_target: str = "",
        extract_sheet: str = "",
        source_id: str = "",
        date_field: str = "",
        start_date: date | None = None,
        end_date: date | None = None,
        refresh_cache: bool = False,
    ) -> list[tuple[str, Record | None]]:
        mode = self.query_source_mode(source, direct)
        operation = {"direct": "批量数据源查询", "extract": "批量提取表查询"}.get(mode, "批量汇总库查询")
        queries = list(dict.fromkeys(value.strip() for value in values if value.strip()))
        if not queries:
            raise ValueError("没有可查询的内容")
        self._log(operation, "INFO", f"开始批量查询：{len(queries)} 个值，字段“{field}”")
        if mode == "extract":
            cache_id = "extract-table"
            cache_target = f"{extract_target}#{extract_sheet}"
            if refresh_cache or not self.cache.has(cache_id) or self.cache._meta(cache_id, "target") != cache_target:
                records = self.read_extract_table(extract_target, extract_sheet)
                self.cache.replace(
                    cache_id,
                    records,
                    schema_field_names(self.store.get("extract_column_schema", [])) or None,
                    cache_target,
                )
            else:
                records = self.cache.load(cache_id)
        elif mode == "direct":
            sources = [item for item in self.store.load_sources() if item.enabled]
            if source_id.strip():
                sources = [item for item in sources if item.id == source_id.strip()]
            records = self.cached_records([item.id for item in sources], refresh=refresh_cache)
        else:
            records = self.database.all_records()
        if date_field.strip():
            if start_date is None or end_date is None:
                raise ValueError("日期范围不完整")
            if start_date > end_date:
                raise ValueError("开始日期不能晚于结束日期")
            before = len(records)
            filtered: list[Record] = []
            valid_date_count = 0
            for record in records:
                current = parse_date(self._field_value(record, date_field))
                if current is not None:
                    valid_date_count += 1
                if current is not None and start_date <= current <= end_date:
                    filtered.append(record)
            records = filtered
            self._log(
                operation,
                "INFO",
                f"日期限制：字段“{date_field}”，{start_date.isoformat()} 至 {end_date.isoformat()}，保留 {len(records)}/{before} 行",
            )
            if before and valid_date_count == 0:
                raise ValueError(
                    f"日期字段“{date_field}”在读取的 {before} 行中没有可识别日期。"
                    "请从下拉框选择当前数据表实际使用的日期列。"
                )
        canonical = self._field_kind(field)
        results: list[tuple[str, Record | None]] = []
        matched_inputs = 0
        indexed_records: list[tuple[Record, list[str]]] = []
        exact_index: dict[str, list[Record]] = {}
        for record in records:
            haystacks = [
                self._match_value(canonical, value)
                for value in self._query_values(record, field)
            ]
            haystacks = self._unique_headers([value for value in haystacks if value])
            indexed_records.append((record, haystacks))
            if exact:
                for haystack in haystacks:
                    exact_index.setdefault(haystack, []).append(record)
        for query_value in queries:
            needle = self._match_value(canonical, query_value)
            if exact:
                matches = exact_index.get(needle, [])
            else:
                matches = [
                    record for record, haystacks in indexed_records
                    if any(needle in haystack for haystack in haystacks)
                ]
            if matches:
                matched_inputs += 1
                results.extend((query_value, record) for record in matches)
            else:
                results.append((query_value, None))
        self._log(
            operation,
            "INFO",
            f"批量查询完成：输入 {len(queries)} 个，找到 {matched_inputs} 个，未找到 {len(queries) - matched_inputs} 个",
        )
        return results

    @staticmethod
    def _match_value(field: str, value: str) -> str:
        text = unicodedata.normalize("NFKC", str(value)).strip().casefold().lstrip("'")
        text = re.sub(r"[\u200b-\u200f\u2060\ufeff]", "", text)
        if field.strip().casefold() not in {"号码", "手机号", "手机号码", "电话", "联系电话", "phone", "number"}:
            return re.sub(r"\s+", "", text)
        compact = re.sub(r"[\s,，\-()（）]", "", text)
        if re.fullmatch(r"[+-]?\d+(?:\.0+)?", compact):
            return compact.split(".", 1)[0].lstrip("+")
        if re.fullmatch(r"[+-]?(?:\d+\.?\d*|\.\d+)[eE][+-]?\d+", compact):
            try:
                return format(Decimal(compact), "f").split(".", 1)[0].lstrip("+")
            except InvalidOperation:
                pass
        digits = re.sub(r"\D", "", compact)
        return digits or compact

    def extract(
        self,
        date_field: str,
        start: date,
        end: date,
        output: str | Path,
        dedup_fields: list[str],
        direct: bool = False,
        output_sheet_name: str = "提取结果",
        destination_type: str = "local",
        google_output_url: str = "",
        source_id: str = "",
    ) -> dict[str, int]:
        operation = "时间提取"
        if start > end:
            raise ValueError("开始日期不能晚于结束日期")
        records = self.read_sources(operation, source_id) if direct else self.database.all_records()
        existing_keys = self._existing_extract_keys(
            destination_type,
            output,
            output_sheet_name,
            google_output_url,
            dedup_fields,
            operation,
        )
        selected: list[tuple[Record, str]] = []
        skipped_invalid = 0
        skipped_duplicate = 0
        for record in records:
            current = parse_date(self._field_value(record, date_field))
            if current is None:
                skipped_invalid += 1
                continue
            if not start <= current <= end:
                continue
            key = self._dedup_key(record, dedup_fields)
            keys = self._dedup_keys(record, dedup_fields)
            if keys & existing_keys or any(self.store.was_extracted(item) for item in keys):
                skipped_duplicate += 1
                continue
            selected.append((record, key))
            existing_keys.update(keys)
        if records and skipped_invalid == len(records):
            raise ValueError(
                f"日期字段“{date_field}”在读取的 {len(records)} 行中没有可识别日期。"
                "请在“日期字段”中选择数据源实际使用的日期列。"
            )
        selected_records = [record for record, _ in selected]
        if destination_type == "google":
            if not google_output_url.strip():
                raise ValueError("请填写目标 Google 表格链接")
            self._write_google_sheet(
                google_output_url,
                output_sheet_name,
                selected_records,
                operation,
                prefer_record_headers=direct,
            )
            destination_label = f"{google_output_url}#{output_sheet_name}"
        else:
            destination = Path(output)
            destination.parent.mkdir(parents=True, exist_ok=True)
            headers = self._signature_headers(
                self._preferred_export_headers(selected_records, prefer_record_headers=direct)
            )
            self._write_xlsx(
                destination,
                selected_records,
                output_sheet_name,
                headers,
                self.store.get("field_aliases", {}),
                self._signature_header(),
                self._signature_value(),
            )
            destination_label = str(destination)
        self.store.mark_extracted([key for _, key in selected], destination_label)
        self._log(
            operation,
            "INFO",
            f"提取完成：写入 {len(selected)} 行，排除重复 {skipped_duplicate} 行，日期无效 {skipped_invalid} 行",
        )
        return {"written": len(selected), "duplicates": skipped_duplicate, "invalid_dates": skipped_invalid}

    def _dedup_parts(self, record: Record, dedup_fields: list[str]) -> list[str]:
        parts = [self._field_value(record, field) for field in dedup_fields]
        if not any(parts):
            parts = [record.spreadsheet_id, record.sheet_name, str(record.row_number), record.row_hash]
        return parts

    def _legacy_dedup_key(self, record: Record, dedup_fields: list[str]) -> str:
        return hashlib.sha256(
            json.dumps(self._dedup_parts(record, dedup_fields), ensure_ascii=False).encode("utf-8")
        ).hexdigest()

    def _dedup_key(self, record: Record, dedup_fields: list[str]) -> str:
        parts = self._dedup_parts(record, dedup_fields)
        normalized = [self._match_value(self._field_kind(field), value) for field, value in zip(dedup_fields, parts)]
        if len(normalized) != len(parts):
            normalized = [str(value).strip() for value in parts]
        return hashlib.sha256(json.dumps(normalized, ensure_ascii=False).encode("utf-8")).hexdigest()

    def _dedup_keys(self, record: Record, dedup_fields: list[str]) -> set[str]:
        return {self._dedup_key(record, dedup_fields), self._legacy_dedup_key(record, dedup_fields)}

    def _existing_extract_keys(
        self,
        destination_type: str,
        output: str | Path,
        output_sheet_name: str,
        google_output_url: str,
        dedup_fields: list[str],
        operation: str,
    ) -> set[str]:
        if destination_type == "google":
            if not google_output_url.strip():
                return set()
            return self._existing_google_extract_keys(
                google_output_url,
                output_sheet_name,
                dedup_fields,
                operation,
            )
        return self._existing_xlsx_extract_keys(Path(output), output_sheet_name, dedup_fields)

    def _records_from_sheet_values(
        self,
        headers: list[str],
        rows: list[list[object]],
        sheet_name: str,
    ) -> list[Record]:
        clean_headers = [str(header or "").strip() for header in headers]
        records: list[Record] = []
        for index, row in enumerate(rows, start=2):
            values: dict[str, str] = {}
            for column, header in enumerate(clean_headers):
                if not header:
                    continue
                value = row[column] if column < len(row) else ""
                values[header] = str(value or "").strip()
            if not any(values.values()):
                continue
            source_name = values.get("来源") or sheet_name
            records.append(
                Record(
                    "extract-output",
                    "提取目标表",
                    "",
                    str(source_name),
                    index,
                    values,
                    "",
                )
            )
        return records

    def _keys_from_existing_rows(
        self,
        headers: list[str],
        rows: list[list[object]],
        sheet_name: str,
        dedup_fields: list[str],
    ) -> set[str]:
        return {
            self._dedup_key(record, dedup_fields)
            for record in self._records_from_sheet_values(headers, rows, sheet_name)
        }

    def _existing_xlsx_extract_keys(
        self,
        path: Path,
        sheet_name: str,
        dedup_fields: list[str],
    ) -> set[str]:
        if not path.exists() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            return set()
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            safe_name = self._safe_sheet_name(sheet_name)
            title = safe_name if safe_name in workbook.sheetnames else workbook.sheetnames[0]
            sheet = workbook[title]
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 2:
                return set()
            headers = [str(cell or "").strip() for cell in rows[0]]
            while headers and not headers[-1]:
                headers.pop()
            return self._keys_from_existing_rows(headers, [list(row) for row in rows[1:]], title, dedup_fields)
        finally:
            workbook.close()

    def _existing_google_extract_keys(
        self,
        url: str,
        sheet_name: str,
        dedup_fields: list[str],
        operation: str,
    ) -> set[str]:
        sid = spreadsheet_id(url)
        if not sid:
            raise ValueError("无法识别目标 Google 表格链接")
        credential = str(self.store.get("credential_path", "")).strip()
        if not credential:
            credential = next(
                (source.credential_path for source in self.store.load_sources() if source.credential_path),
                "",
            )
        if not credential:
            raise ValueError("写入 Google 表格需要在“设置”中配置服务账号 JSON")
        logger = lambda level, message: self._log(operation, level, message)
        client = SourceReader._gspread_client(credential)
        http = client.http_client
        metadata = google_retry(
            lambda: http.fetch_sheet_metadata(
                sid, params={"fields": "sheets.properties(sheetId,title)"}
            ),
            logger,
        )
        target_name = sheet_name.strip() or "提取结果"
        sheet_names = {item["properties"]["title"] for item in metadata.get("sheets", [])}
        if target_name not in sheet_names:
            return set()
        escaped_name = target_name.replace("'", "''")
        response = google_retry(
            lambda: http.values_get(
                sid,
                f"'{escaped_name}'!A:ZZZ",
                params={"majorDimension": "ROWS", "valueRenderOption": "FORMATTED_VALUE"},
            ),
            logger,
        )
        values = response.get("values") or []
        if len(values) < 2:
            return set()
        headers = [str(cell or "").strip() for cell in values[0]]
        while headers and not headers[-1]:
            headers.pop()
        return self._keys_from_existing_rows(headers, values[1:], target_name, dedup_fields)

    def _preferred_export_headers(self, records: list[Record], prefer_record_headers: bool = False) -> list[str]:
        if self.store.get("extract_column_schema_enabled", False):
            configured = schema_field_names(self.store.get("extract_column_schema", []))
            if configured:
                return ["来源", *[name for name in configured if name != "来源"]]
        if prefer_record_headers:
            headers, _ = self._export_rows(records)
            return headers
        if self.store.get("column_schema_enabled", False):
            configured = schema_field_names(self.store.get("column_schema", []))
            if configured:
                return ["来源", *[name for name in configured if name != "来源"]]
        headers, _ = self._export_rows(records)
        return headers

    def _signature_header(self) -> str:
        if not self.store.get("extract_signature_enabled", False):
            return ""
        return str(self.store.get("extract_signature_header", "签字") or "签字").strip()

    def _signature_value(self) -> str:
        if not self.store.get("extract_signature_enabled", False):
            return ""
        return str(self.store.get("extract_signature_value", "") or "").strip()

    def _signature_headers(self, headers: list[str]) -> list[str]:
        if not self.store.get("extract_signature_enabled", False):
            return headers
        signature = self._signature_header()
        if not signature:
            return headers
        result = list(headers)
        if signature in result:
            return result
        column = str(self.store.get("extract_signature_column", "") or "").strip()
        if column:
            index = column_index(column)
            while len(result) <= index:
                result.append("")
            if not result[index]:
                result[index] = signature
            else:
                result.append(signature)
            return result
        result.append(signature)
        return result

    @classmethod
    def _canonical_header(cls, header: str, aliases: dict[str, list[str]]) -> str:
        clean = str(header).strip()
        folded = clean.casefold()
        if folded == "来源":
            return "来源"
        for canonical, candidates in aliases.items():
            if folded in {
                str(name).strip().casefold()
                for name in [canonical, *candidates]
            }:
                return canonical
        semantic = [
            standard
            for standard, hints in cls.SEMANTIC_FIELD_HINTS.items()
            if any(hint in folded for hint in hints)
        ]
        if len(semantic) == 1:
            return semantic[0]
        return clean

    @classmethod
    def _headers_compatible(
        cls,
        existing: list[str],
        required: list[str],
        aliases: dict[str, list[str]],
    ) -> bool:
        return Counter(cls._canonical_header(name, aliases) for name in existing) == Counter(
            cls._canonical_header(name, aliases) for name in required
        )

    @classmethod
    def _rows_for_headers(
        cls,
        records: list[Record],
        headers: list[str],
        aliases: dict[str, list[str]],
        signature_header: str = "",
        signature_value: str = "",
    ) -> list[list[object]]:
        signature_folded = str(signature_header or "").strip().casefold()
        rows: list[list[object]] = []
        for record in records:
            row: list[object] = []
            for header in headers:
                if signature_folded and str(header).strip().casefold() == signature_folded:
                    row.append(signature_value)
                    continue
                canonical = cls._canonical_header(header, aliases)
                if canonical == "来源":
                    row.append(record.sheet_name)
                else:
                    value = record.values.get(canonical, record.values.get(header, ""))
                    if not str(value).strip() and canonical in cls.SEMANTIC_FIELD_HINTS:
                        matches = [
                            candidate
                            for stored, candidate in record.values.items()
                            if str(candidate).strip()
                            and any(hint in str(stored).strip().casefold() for hint in cls.SEMANTIC_FIELD_HINTS[canonical])
                        ]
                        if len(matches) == 1:
                            value = matches[0]
                    row.append(value)
            rows.append(row)
        return rows

    @staticmethod
    def _export_rows(records: list[Record]) -> tuple[list[str], list[list[object]]]:
        fields: list[str] = []
        for record in records:
            for key in record.values:
                if key != "来源" and key not in fields:
                    fields.append(key)
        headers = ["来源", *fields]
        rows = [[record.sheet_name, *[record.values.get(key, "") for key in fields]] for record in records]
        return headers, rows

    @staticmethod
    def _safe_sheet_name(sheet_name: str) -> str:
        safe_name = "".join("_" if char in "[]:*?/\\" else char for char in sheet_name.strip())[:31]
        return safe_name or "提取结果"

    @classmethod
    def _write_xlsx(
        cls,
        path: Path,
        records: list[Record],
        sheet_name: str = "提取结果",
        headers: list[str] | None = None,
        aliases: dict[str, list[str]] | None = None,
        signature_header: str = "",
        signature_value: str = "",
    ) -> None:
        safe_name = cls._safe_sheet_name(sheet_name)
        if path.exists():
            workbook = openpyxl.load_workbook(path)
            sheet = workbook[safe_name] if safe_name in workbook.sheetnames else workbook.create_sheet(safe_name)
        else:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = safe_name
        export_headers, default_rows = cls._export_rows(records)
        existing_header = [
            str(cell.value or "").strip()
            for cell in sheet[1]
        ] if sheet.max_row >= 1 else []
        while existing_header and not existing_header[-1]:
            existing_header.pop()
        headers = existing_header or headers or export_headers
        rows = cls._rows_for_headers(
            records,
            headers,
            aliases or {},
            signature_header,
            signature_value,
        ) if headers != export_headers or aliases or signature_header else default_rows
        if not existing_header:
            for column, header in enumerate(headers, start=1):
                sheet.cell(row=1, column=column, value=header)
        if rows:
            sheet.insert_rows(2, amount=len(rows))
            for row_index, row in enumerate(rows, start=2):
                for column, value in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=column, value=value)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
        workbook.save(path)

    def _write_google_sheet(
        self,
        url: str,
        sheet_name: str,
        records: list[Record],
        operation: str,
        prefer_record_headers: bool = False,
    ) -> None:
        sid = spreadsheet_id(url)
        if not sid:
            raise ValueError("无法识别目标 Google 表格链接")
        credential = str(self.store.get("credential_path", "")).strip()
        if not credential:
            credential = next(
                (source.credential_path for source in self.store.load_sources() if source.credential_path),
                "",
            )
        if not credential:
            raise ValueError("写入 Google 表格需要在“设置”中配置服务账号 JSON")
        logger = lambda level, message: self._log(operation, level, message)
        client = SourceReader._gspread_client(credential)
        http = client.http_client
        metadata = google_retry(
            lambda: http.fetch_sheet_metadata(
                sid, params={"fields": "sheets.properties(sheetId,title)"}
            ),
            logger,
        )
        target_name = sheet_name.strip() or "提取结果"
        aliases = self.store.get("field_aliases", {})
        required_headers = self._signature_headers(self._preferred_export_headers(records, prefer_record_headers))
        sheet_names = {item["properties"]["title"] for item in metadata.get("sheets", [])}
        created = target_name not in sheet_names
        if created:
            google_retry(
                lambda: http.batch_update(
                    sid,
                    {
                        "requests": [{
                            "addSheet": {
                                "properties": {
                                    "title": target_name,
                                    "gridProperties": {
                                        "rowCount": max(1000, len(records) + 10),
                                        "columnCount": max(20, len(required_headers) + 2),
                                    },
                                }
                            }
                        }]
                    },
                ),
                logger,
            )
            existing_header = []
        else:
            escaped_name = target_name.replace("'", "''")
            header_response = google_retry(
                lambda: http.values_get(
                    sid,
                    f"'{escaped_name}'!1:1",
                    params={"majorDimension": "ROWS", "valueRenderOption": "FORMATTED_VALUE"},
                ),
                logger,
            )
            existing_header = (header_response.get("values") or [[]])[0]
            while existing_header and not str(existing_header[-1]).strip():
                existing_header.pop()
        headers = existing_header or required_headers
        headers = self._signature_headers(headers) if existing_header else headers
        rows = self._rows_for_headers(records, headers, aliases, self._signature_header(), self._signature_value())
        payload = ([headers] if not existing_header else []) + rows
        if not payload:
            return
        escaped_name = target_name.replace("'", "''")
        if existing_header:
            if rows:
                sheet_id = next(
                    item["properties"]["sheetId"]
                    for item in metadata.get("sheets", [])
                    if item["properties"]["title"] == target_name
                )
                google_retry(
                    lambda: http.batch_update(
                        sid,
                        {
                            "requests": [{
                                "insertDimension": {
                                    "range": {
                                        "sheetId": sheet_id,
                                        "dimension": "ROWS",
                                        "startIndex": 1,
                                        "endIndex": 1 + len(rows),
                                    },
                                    "inheritFromBefore": False,
                                }
                            }]
                        },
                    ),
                    logger,
                )
            for start in range(0, len(rows), 5000):
                chunk = rows[start:start + 5000]
                google_retry(
                    lambda chunk=chunk, start=start: http.values_update(
                        sid,
                        f"'{escaped_name}'!A{2 + start}",
                        params={"valueInputOption": "RAW"},
                        body={"majorDimension": "ROWS", "values": chunk},
                    ),
                    logger,
                )
                self._log(operation, "INFO", f"已插入 Google 工作表“{target_name}”第 2 行：{min(start + len(chunk), len(rows))}/{len(rows)} 行")
            return
        for start in range(0, len(payload), 5000):
            chunk = payload[start:start + 5000]
            google_retry(
                lambda chunk=chunk, start=start: http.values_update(
                    sid,
                    f"'{escaped_name}'!A{1 + start}",
                    params={"valueInputOption": "RAW"},
                    body={"majorDimension": "ROWS", "values": chunk},
                ),
                logger,
            )
            self._log(operation, "INFO", f"已写入 Google 工作表“{target_name}”：{min(start + len(chunk), len(payload))}/{len(payload)} 行")
